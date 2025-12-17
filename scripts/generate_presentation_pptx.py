"""Generate a PowerPoint presentation from presentation/presentation.md.

The rubric requires a presentation artifact committed to the repo.
This script produces a basic .pptx that keeps Unicode (RU/EN) text.

Usage:
  python scripts/generate_presentation_pptx.py

Output:
  presentation/presentation.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "presentation" / "presentation.md"
OUT = ROOT / "presentation" / "presentation.pptx"


def _parse_markdown(md: str) -> list[tuple[str, list[str]]]:
    """Parse markdown into (title, bullets) slides.

    Heuristic:
    - '# ' starts a new slide title
    - '- ' lines become bullets
    - other non-empty lines become bullets too
    """

    slides: list[tuple[str, list[str]]] = []
    current_title: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_title, current_lines
        if current_title is None:
            return
        bullets = [line.strip() for line in current_lines if line.strip()]
        slides.append((current_title.strip(), bullets))
        current_title = None
        current_lines = []

    for raw in md.splitlines():
        line = raw.rstrip()
        if line.startswith("# "):
            flush()
            current_title = line[2:].strip()
            continue

        if current_title is None:
            continue

        if line.startswith("- "):
            current_lines.append(line[2:].strip())
        else:
            current_lines.append(line.strip())

    flush()
    return slides


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Missing source file: {SRC}")

    md = SRC.read_text(encoding="utf-8")
    slides_data = _parse_markdown(md)

    prs = Presentation()

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Weather EL/ELT Project"
    slide.placeholders[1].text = "Architecture • Data Quality • Analytics"

    # Content slides
    bullet_layout = prs.slide_layouts[1]
    for title, bullets in slides_data:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = title
        body = slide.shapes.placeholders[1].text_frame
        body.clear()

        for i, bullet in enumerate(bullets):
            if not bullet:
                continue
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = bullet
            p.level = 0

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote: {OUT}")


if __name__ == "__main__":
    main()
